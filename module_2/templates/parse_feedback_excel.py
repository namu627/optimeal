"""
parse_feedback_excel.py
=======================
영양사가 작성·반환한 Excel → apply_feedback() → scaling_coefficients.csv 업데이트.

CLI:
    python parse_feedback_excel.py feedback_R01_201_제육볶음_20260504.xlsx
    python parse_feedback_excel.py feedback.xlsx --dry-run   # CSV 수정 없이 결과 출력

워크플로우:
    1. generate_feedback_excel.py로 생성한 xlsx를 영양사에게 발송
    2. 영양사가 '영양사 평가' 시트 작성 후 반환
    3. 이 스크립트 실행 → apply_feedback() → CSV 추가 → JSON 로그 저장

출력:
    - 콘솔: 피드백 요약 (루브릭 점수, 재료별 b 업데이트 결과)
    - feedback_log_R{회차}_{레시피ID}_{날짜}.json: 감사 추적 로그 (삭제 금지)
    - scaling_coefficients.csv: nutritionist_feedback 행 추가 (기존 행 유지)

기준 문서: ADR-002 v3, ADR-003, feedback_rubric_v1.0.md, feedback_json_schema.md
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

# src/engine 경로 추가
_ENGINE_DIR  = Path(__file__).parent.parent / "src" / "engine"
_TMPL_DIR    = Path(__file__).parent
for _p in [str(_ENGINE_DIR), str(_TMPL_DIR)]:
    if _p not in sys.path:
        sys.path.insert(0, _p)

from feedback_update import (  # noqa: E402
    FeedbackResult,
    apply_feedback,
    compute_rubric_avg,
    get_overall_rating,
)
from lookup import build_lookup_table  # noqa: E402
from fallback import get_scaling_params  # noqa: E402

# generate_feedback_excel.py 레이아웃 상수 — 두 파일이 항상 동기화되어야 함
from generate_feedback_excel import (  # noqa: E402
    S1_META,
    S1_ING_START_ROW,
    S1_COL_NAME   as S1_COL_NAME,
    S1_COL_CAT    as S1_COL_CAT,
    S1_COL_BASE   as S1_COL_BASE,
    S1_COL_SCALED as S1_COL_SCALED,
    S2_EVAL_NAME_ROW,
    S2_EVAL_ID_ROW,
    S2_EVAL_DATE_ROW,
    S2_TASTE_ROW,
    S2_FEASIBILITY_ROW,
    S2_FIELD_ROW,
    S2_ING_START_ROW,
    S2_COL_ADJ,
    S2_COL_NOTE,
    S2_COMMENT_OFFSET,
    # 통합 시트 앵커 상수
    ANCHOR_SCALING,
    ANCHOR_EVAL_INFO,
    ANCHOR_RUBRIC,
    ANCHOR_ING_ADJ,
    ANCHOR_COMMENT,
)

_CSV_PATH = _ENGINE_DIR / "scaling_coefficients.csv"


# ─────────────────────────────────────────────────────────────────
# Sheet 1 읽기
# ─────────────────────────────────────────────────────────────────

def _read_meta(ws1, key: str):
    """S1_META 상수를 이용해 Sheet1 메타 값을 읽는다."""
    row, _ = S1_META[key]
    return ws1.cell(row=row, column=2).value


def _parse_recipe_sheet(ws1) -> dict:
    """Sheet 1에서 레시피 메타·재료 데이터를 읽어 dict로 반환한다."""
    recipe_name    = str(_read_meta(ws1, "recipe_name") or "")
    round_raw      = str(_read_meta(ws1, "feedback_round") or "1")   # "1회차"
    feedback_round = int(round_raw.replace("회차", "").strip())
    _rid_raw = _read_meta(ws1, "recipe_id") or "0"
    try:
        recipe_id: int | str = int(_rid_raw)
    except (ValueError, TypeError):
        recipe_id = str(_rid_raw)  # 문자열 ID 허용 (예: "SAMPLE_001")

    serving_raw    = str(_read_meta(ws1, "target_serving") or "1인분 → 100인분")  # "1인분 → 100인분"
    target_serving_n = int(serving_raw.split("→")[-1].replace("인분", "").strip())

    method_raw = str(_read_meta(ws1, "cooking_method") or "볶음 (dry_heat)")  # "볶음 (dry_heat)"
    group_type = method_raw.split("(")[-1].rstrip(")").strip()

    # 재료 행 읽기 (이름이 빈 셀을 만나면 종료)
    ingredients: list[dict] = []
    for row in range(S1_ING_START_ROW, ws1.max_row + 1):
        name = ws1.cell(row=row, column=S1_COL_NAME).value
        if not name:
            break
        ingredients.append({
            "name":          str(name),
            "category":      str(ws1.cell(row=row, column=S1_COL_CAT).value or ""),
            "base_amount_g": float(ws1.cell(row=row, column=S1_COL_BASE).value or 0),
            "scaled_amount_g": float(ws1.cell(row=row, column=S1_COL_SCALED).value or 0),
        })

    return {
        "recipe_id":       recipe_id,
        "recipe_name":     recipe_name,
        "group_type":      group_type,
        "feedback_round":  feedback_round,
        "target_serving_n": target_serving_n,
        "ingredients":     ingredients,
    }


# ─────────────────────────────────────────────────────────────────
# Sheet 2 읽기
# ─────────────────────────────────────────────────────────────────

def _parse_evaluation_sheet(ws2, recipe: dict) -> tuple[dict, dict]:
    """
    Sheet 2에서 평가자 정보, 루브릭 점수, 재료별 조정값을 읽는다.

    Returns:
        evaluator: {name, id, date, taste, feasibility, field, comment}
        adjustments: {재료명: delta_g_float}  — 0g는 포함하지 않음
    """
    def _val(row, col=2):
        return ws2.cell(row=row, column=col).value

    evaluator = {
        "name":        str(_val(S2_EVAL_NAME_ROW)   or ""),
        "id":          str(_val(S2_EVAL_ID_ROW)     or ""),
        "date":        str(_val(S2_EVAL_DATE_ROW)   or str(date.today())),
        "taste":       _val(S2_TASTE_ROW),
        "feasibility": _val(S2_FEASIBILITY_ROW),
        "field":       _val(S2_FIELD_ROW),
    }

    # 재료별 조정값 (0 및 빈칸 제외)
    adjustments: dict[str, float] = {}
    for i, ing in enumerate(recipe["ingredients"]):
        row     = S2_ING_START_ROW + i
        adj_raw = ws2.cell(row=row, column=S2_COL_ADJ).value
        if adj_raw is not None and str(adj_raw).strip() not in ("", "0"):
            try:
                val = float(adj_raw)
                if val != 0.0:
                    adjustments[ing["name"]] = val
            except (ValueError, TypeError):
                pass

    # 종합 의견 (로그용)
    comment_row        = S2_ING_START_ROW + len(recipe["ingredients"]) + S2_COMMENT_OFFSET + 1
    evaluator["comment"] = str(ws2.cell(row=comment_row, column=1).value or "")

    return evaluator, adjustments


# ─────────────────────────────────────────────────────────────────
# 루브릭 유효성 검사
# ─────────────────────────────────────────────────────────────────

def _validate_rubric(evaluator: dict) -> None:
    """루브릭 점수가 1~5 정수인지 확인한다. 위반 시 ValueError."""
    errors: list[str] = []
    for key, label in [
        ("taste",       "맛·간 적절성"),
        ("feasibility", "조리 실현 가능성"),
        ("field",       "대량 조리 현장 적합성"),
    ]:
        v = evaluator[key]
        if v not in {1, 2, 3, 4, 5}:
            errors.append(f"  '{label}' 점수 오류: {repr(v)}  → 1~5 정수 필요")
    if errors:
        raise ValueError("루브릭 점수 유효성 검사 실패:\n" + "\n".join(errors))


# ─────────────────────────────────────────────────────────────────
# 현재 파라미터 조회 (회차 누적 지원)
# ─────────────────────────────────────────────────────────────────

def _get_current_params(group_type: str, ingredient_category: str) -> tuple[float, float, float, int]:
    """
    CSV에서 현재 (group_type × ingredient_category) 파라미터를 읽는다.

    nutritionist_feedback 행이 있으면 최신 행 우선 (회차 누적 반영).
    없으면 lookup/fallback 결과를 사용한다.

    Returns:
        (a, b, se_b, n_old)
        n_old: 기존 nutritionist_feedback 행 수 (이동 평균 분모)
    """
    df = pd.read_csv(_CSV_PATH)
    fb = df[
        (df["group_type"] == group_type) &
        (df["ingredient_category"] == ingredient_category) &
        (df["estimation_method"] == "nutritionist_feedback")
    ].sort_values("coefficient_id", ascending=False)

    if len(fb) > 0:
        row = fb.iloc[0]
        return float(row["power_law_a"]), float(row["power_law_b"]), float(row["se_b"]), len(fb)

    # MixedLM / fallback 파라미터
    lookup_df = build_lookup_table()
    params    = get_scaling_params(lookup_df, ingredient_category, group_type)
    return params.power_law_a, params.power_law_b, params.se_b, 0


# ─────────────────────────────────────────────────────────────────
# CSV 업데이트 (추가 전용 — 기존 행 수정 금지)
# ─────────────────────────────────────────────────────────────────

def _append_csv_row(
    group_type: str,
    ingredient_category: str,
    a_old: float,
    b_new: float,
    se_b: float,
    notes: str,
) -> None:
    """
    scaling_coefficients.csv에 nutritionist_feedback 행을 추가한다.

    lookup_priority=1 이므로 이후 룩업에서 MixedLM(priority=2)보다 우선 조회된다.
    기존 행은 절대 수정하지 않는다 (ADR-003 감사 추적 요건).
    """
    df     = pd.read_csv(_CSV_PATH)
    new_id = int(df["coefficient_id"].max()) + 1

    new_row = pd.DataFrame([{
        "coefficient_id":      new_id,
        "group_type":          group_type,
        "ingredient_category": ingredient_category,
        "power_law_a":         round(a_old, 4),
        "power_law_b":         round(b_new, 4),
        "se_b":                round(se_b, 4),
        "ci_lower":            "",
        "ci_upper":            "",
        "r_squared":           "",
        "n_obs":               "",
        "estimation_method":   "nutritionist_feedback",
        "lookup_priority":     1,
        "source":              "영양사피드백",
        "notes":               notes,
    }])
    pd.concat([df, new_row], ignore_index=True).to_csv(_CSV_PATH, index=False)


# ─────────────────────────────────────────────────────────────────
# 통합 시트 파싱 (generate_single_file 형식)
# ─────────────────────────────────────────────────────────────────

def _find_anchor_row(ws, anchor: str) -> int:
    """워크시트 A열에서 앵커 문자열로 시작하는 행 번호를 반환한다."""
    for row in ws.iter_rows(min_col=1, max_col=1):
        v = row[0].value
        if v and str(v).startswith(anchor):
            return row[0].row
    raise ValueError(f"앵커 '{anchor}'를 찾을 수 없습니다. (시트: {ws.title})")


def _parse_combined_sheet(ws) -> tuple[dict, dict, dict]:
    """통합 시트에서 레시피 메타, 평가자 정보, 재료 조정값을 파싱한다.

    Returns:
        (recipe, evaluator, adjustments)
        - recipe: _parse_recipe_sheet()와 동일 구조
        - evaluator: {name, id, date, taste, feasibility, field, comment}
        - adjustments: {재료명: delta_g} (0은 제외)
    """
    def _cv(row, col):
        return ws.cell(row=row, column=col).value

    # ── 메타 (고정 행 3-5) ────────────────────────────────
    recipe_name = str(_cv(3, 2) or "")

    serving_raw  = str(_cv(4, 2) or "1인분 → 100인분")
    target_n     = int(serving_raw.split("→")[-1].replace("인분", "").strip())

    rid_raw = _cv(5, 2) or "0"
    try:
        recipe_id: int | str = int(rid_raw)
    except (ValueError, TypeError):
        recipe_id = str(rid_raw)

    method_raw = str(_cv(3, 6) or "볶음 (dry_heat)")   # F열 (col 6)
    group_type = method_raw.split("(")[-1].rstrip(")").strip()

    round_raw      = str(_cv(4, 6) or "1회차")
    feedback_round = int(round_raw.replace("회차", "").strip())

    # ── 스케일링 결과 → 재료 목록 ────────────────────────
    sc_row = _find_anchor_row(ws, ANCHOR_SCALING)
    # 재료 데이터: sc_row+2 부터 이름 빈 셀 만날 때까지
    ingredients: list[dict] = []
    for row in range(sc_row + 2, ws.max_row + 1):
        name = _cv(row, 2)
        if not name:
            break
        ingredients.append({
            "name":            str(name),
            "category":        str(_cv(row, 3) or ""),
            "base_amount_g":   float(_cv(row, 4) or 0),
            "scaled_amount_g": float(_cv(row, 5) or 0),
        })

    recipe = {
        "recipe_id":        recipe_id,
        "recipe_name":      recipe_name,
        "group_type":       group_type,
        "feedback_round":   feedback_round,
        "target_serving_n": target_n,
        "ingredients":      ingredients,
    }

    # ── 평가자 정보 ───────────────────────────────────────
    ei_row = _find_anchor_row(ws, ANCHOR_EVAL_INFO)
    evaluator = {
        "name":  str(_cv(ei_row + 1, 2) or ""),
        "id":    str(_cv(ei_row + 2, 2) or ""),
        "date":  str(_cv(ei_row + 3, 2) or str(date.today())),
    }

    # ── 루브릭 ───────────────────────────────────────────
    rb_row = _find_anchor_row(ws, ANCHOR_RUBRIC)
    evaluator["taste"]       = _cv(rb_row + 2, 2)
    evaluator["feasibility"] = _cv(rb_row + 3, 2)
    evaluator["field"]       = _cv(rb_row + 4, 2)

    # ── 재료별 수정 의견 ─────────────────────────────────
    ia_row    = _find_anchor_row(ws, ANCHOR_ING_ADJ)
    adj_start = ia_row + 2
    adjustments: dict[str, float] = {}
    for i, ing in enumerate(ingredients):
        adj_raw = _cv(adj_start + i, 5)   # E열 = col 5
        if adj_raw is not None and str(adj_raw).strip() not in ("", "0"):
            try:
                val = float(adj_raw)
                if val != 0.0:
                    adjustments[ing["name"]] = val
            except (ValueError, TypeError):
                pass

    # ── 종합 의견 ─────────────────────────────────────────
    cm_row = _find_anchor_row(ws, ANCHOR_COMMENT)
    evaluator["comment"] = str(_cv(cm_row + 1, 1) or "")

    return recipe, evaluator, adjustments


def _apply_and_build_log(
    recipe: dict,
    evaluator: dict,
    adjustments: dict,
    source_label: str,
    dry_run: bool,
) -> dict:
    """apply_feedback 실행 후 감사 추적 로그 dict를 반환한다."""
    _validate_rubric(evaluator)

    rubric_avg     = compute_rubric_avg(
        int(evaluator["taste"]), int(evaluator["feasibility"]), int(evaluator["field"])
    )
    overall_rating = get_overall_rating(rubric_avg)

    ingredient_updates: list[dict] = []
    for ing in recipe["ingredients"]:
        delta_g = adjustments.get(ing["name"], 0.0)
        if delta_g == 0.0:
            continue

        a_old, b_old, se_b, n_old = _get_current_params(
            recipe["group_type"], ing["category"]
        )
        result: FeedbackResult = apply_feedback(
            actual_amount_g  = ing["scaled_amount_g"],
            delta_g_original = delta_g,
            base_amount_g    = ing["base_amount_g"],
            target_serving_n = recipe["target_serving_n"],
            b_old            = b_old,
            n_old            = n_old,
            a_old            = a_old,
            se_b             = se_b,
            delta_g_history  = [],
        )

        notes = (
            f"R{recipe['feedback_round']} {recipe['recipe_name']} | "
            f"{ing['name']}({ing['category']}) | "
            f"평가자:{evaluator['id']} | "
            f"Δ_g:{delta_g:+.1f} | rubric_avg:{rubric_avg}"
        )

        if not dry_run:
            _append_csv_row(
                group_type          = recipe["group_type"],
                ingredient_category = ing["category"],
                a_old               = a_old,
                b_new               = result.b_new,
                se_b                = se_b,
                notes               = notes,
            )

        ingredient_updates.append({
            "ingredient_name":  ing["name"],
            "category":         ing["category"],
            "delta_g_original": result.delta_g_original,
            "delta_g_clipped":  result.delta_g_clipped,
            "b_old":            round(b_old, 4),
            "b_feedback":       round(result.b_feedback, 4),
            "b_new_raw":        round(result.b_new_raw, 4),
            "b_new":            round(result.b_new, 4),
            "is_clipped":       result.is_clipped,
            "n_new":            result.n_new,
        })

    return {
        "source":            source_label,
        "recipe_id":         recipe["recipe_id"],
        "recipe_name":       recipe["recipe_name"],
        "feedback_round":    recipe["feedback_round"],
        "group_type":        recipe["group_type"],
        "evaluator_name":    evaluator["name"],
        "evaluator_id":      evaluator["id"],
        "evaluation_date":   evaluator["date"],
        "rubric_taste":      evaluator["taste"],
        "rubric_feasibility": evaluator["feasibility"],
        "rubric_field":      evaluator["field"],
        "rubric_avg":        rubric_avg,
        "overall_rating":    overall_rating,
        "comment":           evaluator.get("comment", ""),
        "ingredient_updates": ingredient_updates,
        "dry_run":           dry_run,
        "processed_date":    str(date.today()),
    }


def parse_single_file(path: Path, dry_run: bool = False) -> list[dict]:
    """
    통합 Excel 파일 (시트 = 레시피)을 전체 파싱한다.

    generate_single_file()로 생성한 파일을 처리한다.
    각 시트를 _parse_combined_sheet()로 파싱 → apply_feedback 실행.

    Args:
        path:     generate_all_rounds.py가 생성한 통합 .xlsx 파일 경로
        dry_run:  True이면 CSV를 수정하지 않음

    Returns:
        시트별 감사 추적 로그 dict 목록
    """
    wb   = openpyxl.load_workbook(path, data_only=True)
    logs: list[dict] = []

    for ws in wb.worksheets:
        try:
            recipe, evaluator, adjustments = _parse_combined_sheet(ws)
            log = _apply_and_build_log(recipe, evaluator, adjustments,
                                       source_label=f"{path.name}::{ws.title}",
                                       dry_run=dry_run)
            logs.append(log)
        except Exception as e:
            print(f"[경고] 시트 '{ws.title}' 파싱 실패: {e}", file=sys.stderr)

    return logs


# ─────────────────────────────────────────────────────────────────
# 메인 파싱 함수 (기존 2-시트 형식)
# ─────────────────────────────────────────────────────────────────

def parse(excel_path: Path, dry_run: bool = False) -> dict:
    """
    영양사 작성 Excel → apply_feedback() → CSV 업데이트.
    기존 2-시트 형식 (Sheet "레시피" + "영양사 평가") 처리용.

    Args:
        excel_path: 영양사가 작성한 .xlsx 파일 경로
        dry_run: True이면 CSV를 수정하지 않음 (결과 확인용)

    Returns:
        감사 추적 로그 dict (JSON 저장과 동일한 구조)
    """
    wb   = openpyxl.load_workbook(excel_path, data_only=True)
    ws1  = wb["레시피"]
    ws2  = wb["영양사 평가"]

    recipe             = _parse_recipe_sheet(ws1)
    evaluator, adj_map = _parse_evaluation_sheet(ws2, recipe)

    log = _apply_and_build_log(recipe, evaluator, adj_map,
                               source_label=excel_path.name,
                               dry_run=dry_run)
    log["excel_file"] = excel_path.name   # 이전 호환성 유지
    return log


# ─────────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────────

def _print_log_summary(log: dict) -> None:
    """단일 레시피 로그 요약을 콘솔에 출력한다."""
    W = 65
    print(f"\n{'='*W}")
    print(f"  레시피: {log['recipe_name']}  ({log['feedback_round']}회차 / {log['group_type']})")
    print(f"  평가자: {log['evaluator_name']}  ({log['evaluator_id']})  / 날짜: {log['evaluation_date']}")
    print(f"{'─'*W}")
    print(f"  루브릭:  ① {log['rubric_taste']}  ② {log['rubric_feasibility']}  ③ {log['rubric_field']}"
          f"  →  평균 {log['rubric_avg']}  [{log['overall_rating']}]")
    print(f"{'─'*W}")
    updates = log["ingredient_updates"]
    if updates:
        print(f"  재료별 파라미터 업데이트 ({len(updates)}건)")
        for r in updates:
            clip = "  ⚠ 클리핑" if r["is_clipped"] else ""
            print(f"    {r['ingredient_name']:<12s}  b: {r['b_old']:.4f} → {r['b_new']:.4f}"
                  f"  Δ_g={r['delta_g_original']:+.1f}g{clip}")
    else:
        print("  재료 수정 없음 (모든 조정량 0g)")
    if log.get("comment"):
        print(f"  종합 의견: {log['comment'][:80]}{'...' if len(log['comment']) > 80 else ''}")
    print(f"{'='*W}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "영양사 피드백 Excel → apply_feedback → scaling_coefficients.csv 업데이트.\n"
            "통합 파일 (피드백_전체_5회차_*.xlsx) 및 기존 단건 파일 모두 지원."
        )
    )
    parser.add_argument("excel_path", help="영양사가 작성·반환한 .xlsx 파일 경로")
    parser.add_argument("--dry-run", action="store_true",
                        help="CSV 수정 없이 결과만 출력 (검토용)")
    parser.add_argument("--log-dir", default=".",
                        help="JSON 로그 저장 폴더 (기본: 현재 디렉토리)")
    args = parser.parse_args()

    path = Path(args.excel_path)
    if not path.exists():
        print(f"[오류] 파일 없음: {path}", file=sys.stderr)
        sys.exit(1)

    # ── 파일 형식 자동 감지 ──────────────────────────────
    # 통합 파일: 시트 이름이 "레시피"/"영양사 평가"가 아닌 경우 (R1_R1_001_… 형식)
    wb_check = openpyxl.load_workbook(path, read_only=True)
    sheet_names = wb_check.sheetnames
    wb_check.close()

    is_legacy = ("레시피" in sheet_names and "영양사 평가" in sheet_names)

    if is_legacy:
        logs = [parse(path, dry_run=args.dry_run)]
    else:
        logs = parse_single_file(path, dry_run=args.dry_run)

    # ── 콘솔 요약 출력 ──────────────────────────────────
    total_updates = sum(len(lg["ingredient_updates"]) for lg in logs)
    if len(logs) > 1:
        print(f"\n[통합 파일] {len(logs)}개 시트 파싱 완료 / 파라미터 업데이트 {total_updates}건")

    for log in logs:
        _print_log_summary(log)

    if args.dry_run:
        print(f"\n[dry-run 모드]  scaling_coefficients.csv 는 수정되지 않았습니다.")
    else:
        print(f"\n[완료]  scaling_coefficients.csv  →  {total_updates}행 추가됨.")

    # ── JSON 로그 저장 ──────────────────────────────────
    log_dir = Path(args.log_dir)
    for log in logs:
        log_name = (
            f"feedback_log"
            f"_R{log['feedback_round']:02d}"
            f"_{log['recipe_id']}"
            f"_{date.today().strftime('%Y%m%d')}.json"
        )
        log_path = log_dir / log_name
        log_path.write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"[로그 저장]  {log_path.resolve()}")


if __name__ == "__main__":
    main()
